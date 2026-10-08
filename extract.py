"""Extracts per-player data from 2026_Shepherds_Live_AB_Data.xlsx for the report generator."""
import csv
import math
import os
import re
import openpyxl
from collections import defaultdict
import benchmarks as bm

DATA_XLSX = '/mnt/user-data/uploads/2026_Shepherds_Live_AB_Data.xlsx'  # default/fallback for CLI use

PITCH_TYPE_MAP = {
    'FB': 'Fastball', 'FB 1': 'Fastball', 'FB 2': 'Fastball', 'FB S': 'Fastball',
    'SL': 'Slider', 'SL 2': 'Slider', 'Sl': 'Slider',
    'CH': 'Changeup', 'CH 1': 'Changeup',
    'CB': 'Curveball', 'CB 1': 'Curveball', 'CB 2': 'Curveball',
    'CT': 'Cutter', 'SP': 'Splitter',
}
# Same palette as MINI_PITCH_COLORS below (sampled from the client's reference
# pitch-type legend), keyed by Full Swing's coarser labels so every pitcher
# chart in the app -- full report and Trackman-only -- uses one color scheme.
PITCH_COLORS = {
    'Fastball': '#c13d4d', 'Sinker': '#f0a139', 'Slider': '#ede750', 'Sweeper': '#d6b552',
    'Changeup': '#5abb4e', 'Curveball': '#5fcee9', 'Cutter': '#894432', 'Splitter': '#5daaab',
}


def _num(v):
    return v if isinstance(v, (int, float)) else None


def load_totals(xlsx_path=None):
    wb = openpyxl.load_workbook(xlsx_path or DATA_XLSX, data_only=True)
    ws = wb['Totals']

    def read_block(header_row, start_row, end_row):
        hdr = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        rows = []
        for r in range(start_row, end_row + 1):
            name = ws.cell(row=r, column=3).value
            if not name:
                continue
            d = {}
            for i, h in enumerate(hdr, start=1):
                if h is not None:
                    d[h] = ws.cell(row=r, column=i).value
            rows.append(d)
        return rows

    pitching = read_block(3, 4, 10)
    hitting = read_block(16, 17, 25)
    return pitching, hitting


def load_master_rows(xlsx_path=None):
    wb = openpyxl.load_workbook(xlsx_path or DATA_XLSX, data_only=True)
    ws = wb['Master']
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['Pitcher']] is None:
            continue
        rows.append({h: row[i] for h, i in idx.items()})
    return rows


def pitcher_tracking_table(totals_row):
    specs = [
        ('Fastball', 'TFB', 'FB USG', 'FB K%', 'FB S/M%'),
        ('Curveball', 'TCB', 'CB USG', 'CB K%', 'CB S/M%'),
        ('Changeup', 'TCH', 'CH USG', 'CH K%', 'CH S/M%'),
        ('Other', 'TO', 'O USG', 'O K%', 'O S/M%'),
    ]
    out = []
    for label, ncol, usgcol, kcol, smcol in specs:
        n = totals_row.get(ncol) or 0
        if not n:
            continue
        usg = _num(totals_row.get(usgcol)) or 0
        kpct = _num(totals_row.get(kcol)) or 0
        smpct = _num(totals_row.get(smcol)) or 0
        out.append({'type': label, 'n': int(n), 'pct': usg * 100, 'k_pct': kpct * 100, 'whiff_pct': smpct * 100})
    return out


# ---------------------------------------------------------------------------
# Master-tab-only pitcher stats -- no reliance on the Totals sheet's fixed
# Fastball/Curveball/Changeup/Other columns at all. Pitch-type labels come
# straight from each pitch's own "Pitch Type" cell (via PITCH_TYPE_MAP), so
# a pitcher who actually throws Slider/Cutter/Splitter instead of a true
# Curveball is represented correctly -- this is what fixed the Esteban/
# Samples Curveball-vs-Slider mislabeling.
#
# Plate-appearance (BF) boundaries aren't stored explicitly in Master, so
# they're inferred: consecutive rows with the same Pitcher+Batter belong to
# one PA, which ends the moment a pitch's outcome is one of the PA-ending
# values below (or the Pitcher/Batter changes without one, as a fallback
# for any incomplete/edge-case AB).
_PA_END_OUTCOMES = {'BB', 'K', 'bK', 'In Play', 'HBP'}
_SWING_OUTCOMES = {'Foul', 'In Play', 'Yes'}  # 'Yes' = swinging strike/whiff


def _group_plate_appearances(master_rows, pitcher_name):
    pas, cur, cur_key = [], [], None
    for r in master_rows:
        if r.get('Pitcher') != pitcher_name:
            continue
        key = (r.get('Pitcher'), r.get('Batter'))
        if cur_key is not None and key != cur_key and cur:
            pas.append(cur)
            cur = []
        cur_key = key
        cur.append(r)
        if r.get('Sw/Miss/Outcome') in _PA_END_OUTCOMES:
            pas.append(cur)
            cur, cur_key = [], None
    if cur:
        pas.append(cur)
    return pas


def pitcher_master_stats(master_rows, pitcher_name):
    """Everything the report needs about a pitcher, computed straight from
    Master rows: overall rate stats (for the percentile bars) plus a
    per-pitch-type tracking table (usage/K%/whiff%/velo/spin/EV/SO),
    keyed off real per-pitch Pitch Type labels rather than Totals columns.

    NOTE: G / IP / H / WHIP still aren't derivable from Master alone --
    there's no outs-recorded or hit-vs-out-on-balls-in-play data here, so
    those keep coming from the Totals sheet (or the per-player sheets)
    until that data is captured at the point of entry.
    """
    prows = [r for r in master_rows if r.get('Pitcher') == pitcher_name]
    pas = _group_plate_appearances(master_rows, pitcher_name)
    bf = len(pas)

    k = bb = hbp = fps = 0
    strikes = pitches_with_bs = swings = whiffs = 0
    for pa in pas:
        last = pa[-1].get('Sw/Miss/Outcome')
        if last in ('K', 'bK'):
            k += 1
        elif last == 'BB':
            bb += 1
        elif last == 'HBP':
            hbp += 1
        if pa[0].get('Ball/Strike') == 1:
            fps += 1
        for p in pa:
            bs = p.get('Ball/Strike')
            if bs is not None:
                pitches_with_bs += 1
                if bs == 1:
                    strikes += 1
            oc = p.get('Sw/Miss/Outcome')
            if oc in _SWING_OUTCOMES:
                swings += 1
                if oc == 'Yes':
                    whiffs += 1

    overall = {
        'bf': bf, 'k': k, 'bb': bb, 'hbp': hbp,
        'strike_pct': round(100 * strikes / pitches_with_bs, 1) if pitches_with_bs else None,
        'fps_pct': round(100 * fps / bf, 1) if bf else None,
        'k_pct': round(100 * k / bf, 1) if bf else None,
        'bb_pct': round(100 * bb / bf, 1) if bf else None,
        'whiff_pct': round(100 * whiffs / swings, 1) if swings else None,
    }

    by_type = defaultdict(lambda: {'n': 0, 'k': 0, 'swings': 0, 'whiffs': 0,
                                    'velos': [], 'spins': [], 'evs': [], 'so': 0,
                                    'strikes': 0, 'pitches_with_bs': 0})
    total_pitches = len(prows)
    for r in prows:
        label = PITCH_TYPE_MAP.get(r.get('Pitch Type'))
        if not label:
            continue
        d = by_type[label]
        d['n'] += 1
        oc = r.get('Sw/Miss/Outcome')
        if oc in ('K', 'bK'):
            d['k'] += 1
            d['so'] += 1
        if oc in _SWING_OUTCOMES:
            d['swings'] += 1
            if oc == 'Yes':
                d['whiffs'] += 1
        if _num(r.get('Velocity')) is not None:
            d['velos'].append(r['Velocity'])
        if _num(r.get('SpinRate')) is not None:
            d['spins'].append(r['SpinRate'])
        if _num(r.get('ExitSpeed')) is not None:
            d['evs'].append(r['ExitSpeed'])
        # Ball/Strike: 1 = strike, 0 = ball -- per-pitch-type Strike%,
        # same convention as the overall Strike% above.
        bs = r.get('Ball/Strike')
        if bs is not None:
            d['pitches_with_bs'] += 1
            if bs == 1:
                d['strikes'] += 1

    tracking = []
    for label, d in sorted(by_type.items(), key=lambda kv: -kv[1]['n']):
        tracking.append({
            'type': label,
            'n': d['n'],
            'pct': 100 * d['n'] / total_pitches if total_pitches else 0,
            'k_pct': 100 * d['k'] / d['n'] if d['n'] else 0,
            'whiff_pct': 100 * d['whiffs'] / d['swings'] if d['swings'] else 0,
            'strike_pct': round(100 * d['strikes'] / d['pitches_with_bs'], 1) if d['pitches_with_bs'] else None,
            'velo': round(sum(d['velos']) / len(d['velos']), 1) if d['velos'] else None,
            'top_velo': round(max(d['velos']), 1) if d['velos'] else None,
            'spin': round(sum(d['spins']) / len(d['spins'])) if d['spins'] else None,
            'ev': round(sum(d['evs']) / len(d['evs']), 1) if d['evs'] else None,
            'so': d['so'],
        })
    return overall, tracking


def pitcher_spin_ev_so(master_rows, pitcher_name):
    buckets = defaultdict(lambda: {'spins': [], 'evs': [], 'so': 0, 'velos': []})
    for row in master_rows:
        if row.get('Pitcher') != pitcher_name:
            continue
        label = PITCH_TYPE_MAP.get(row.get('Pitch Type'))
        if not label:
            continue
        if _num(row.get('SpinRate')) is not None:
            buckets[label]['spins'].append(row['SpinRate'])
        if _num(row.get('ExitSpeed')) is not None:
            buckets[label]['evs'].append(row['ExitSpeed'])
        if _num(row.get('Velocity')) is not None:
            buckets[label]['velos'].append(row['Velocity'])
        outcome = row.get('Sw/Miss/Outcome') or row.get('Outcome')
        if outcome in ('K', 'bK'):
            buckets[label]['so'] += 1
    result = {}
    for label, d in buckets.items():
        result[label] = {
            'spin': round(sum(d['spins']) / len(d['spins'])) if d['spins'] else None,
            'ev': round(sum(d['evs']) / len(d['evs']), 1) if d['evs'] else None,
            'so': d['so'],
            'velo': round(sum(d['velos']) / len(d['velos']), 1) if d['velos'] else None,
            'top_velo': round(max(d['velos']), 1) if d['velos'] else None,
            'n': len(d['velos']),
        }
    return result


def hitter_spray_points(master_rows, hitter_name):
    pts = []
    for row in master_rows:
        if row.get('Batter') != hitter_name:
            continue
        if _num(row.get('Direction')) is not None and _num(row.get('Distance')) is not None:
            pts.append({
                'direction': row['Direction'], 'distance': row['Distance'],
                'ev': _num(row.get('ExitSpeed')), 'la': _num(row.get('Angle')),
            })
    return pts


def hitter_derived_metrics(master_rows, hitter_name):
    balls = []
    for row in master_rows:
        if row.get('Batter') != hitter_name:
            continue
        ev = _num(row.get('ExitSpeed'))
        if ev is None:
            continue
        balls.append({'ev': ev, 'dist': _num(row.get('Distance')), 'angle': _num(row.get('Angle')), 'squp': _num(row.get('SquaredUp'))})
    if not balls:
        return {}
    top_half = sorted(balls, key=lambda b: b['ev'], reverse=True)[:max(1, len(balls) // 2)]
    dists = [b['dist'] for b in top_half if b['dist'] is not None]
    angles = [b['angle'] for b in balls if b['angle'] is not None]
    squps = [b['squp'] for b in balls if b['squp'] is not None]
    return {
        'top50_dist': round(sum(dists) / len(dists), 1) if dists else None,
        'avg_launch_angle': round(sum(angles) / len(angles), 1) if angles else None,
        'squared_up_pct': round(100 * sum(squps) / len(squps), 1) if squps else None,
    }


def pitcher_velo_ev_metrics(master_rows, pitcher_name):
    """Fastball Velo and Avg Exit Velo Allowed aren't on the Totals sheet at
    all — they only exist as raw per-pitch values on the Master tab — so
    they're computed here the same way pitcher_spin_ev_so derives its
    per-pitch-type numbers, just rolled up across all of this pitcher's
    pitches instead of split by type."""
    fb_velos, all_evs = [], []
    for row in master_rows:
        if row.get('Pitcher') != pitcher_name:
            continue
        label = PITCH_TYPE_MAP.get(row.get('Pitch Type'))
        if label == 'Fastball' and _num(row.get('Velocity')) is not None:
            fb_velos.append(row['Velocity'])
        if _num(row.get('ExitSpeed')) is not None:
            all_evs.append(row['ExitSpeed'])
    return {
        'fb_velo': round(sum(fb_velos) / len(fb_velos), 1) if fb_velos else None,
        'avg_ev_allowed': round(sum(all_evs) / len(all_evs), 1) if all_evs else None,
    }


def build_pitcher_percentiles(totals_row, velo_ev=None):
    def pct(key, val):
        return bm.interpolate(val, bm.PITCHER[key]) if val is not None else None

    velo_ev = velo_ev or {}
    fb_velo = _num(velo_ev.get('fb_velo'))
    avg_ev_allowed = _num(velo_ev.get('avg_ev_allowed'))

    strike_pct = _num(totals_row.get('S%'))
    strike_pct = strike_pct * 100 if strike_pct is not None else None
    fps_pct = _num(totals_row.get('FPS%'))
    fps_pct = fps_pct * 100 if fps_pct is not None else None
    k_pct = _num(totals_row.get('K%'))
    k_pct = k_pct * 100 if k_pct is not None else None
    bb_pct = _num(totals_row.get('BB%'))
    bb_pct = bb_pct * 100 if bb_pct is not None else None
    whiff_pct = _num(totals_row.get('SW/M%'))
    whiff_pct = whiff_pct * 100 if whiff_pct is not None else None

    rows = [
        {'label': 'FB Velo', 'value': fb_velo, 'pctl': pct('fb_velo', fb_velo), 'unit': ''},
        {'label': 'Avg EV Allowed', 'value': avg_ev_allowed, 'pctl': pct('avg_ev_allowed', avg_ev_allowed), 'unit': ''},
        {'label': 'Strike%', 'value': strike_pct, 'pctl': pct('strike_pct', strike_pct), 'unit': '%'},
        {'label': 'FPS%', 'value': fps_pct, 'pctl': pct('fps_pct', fps_pct), 'unit': '%'},
        {'label': 'K%', 'value': k_pct, 'pctl': pct('k_pct', k_pct), 'unit': '%'},
        {'label': 'BB%', 'value': bb_pct, 'pctl': pct('bb_pct', bb_pct), 'unit': '%'},
        {'label': 'Whiff%', 'value': whiff_pct, 'pctl': pct('whiff_pct', whiff_pct), 'unit': '%'},
    ]
    return [r for r in rows if r['value'] is not None]


def build_pitcher_percentiles_master(overall, velo_ev=None):
    """Same output shape as build_pitcher_percentiles(), but Strike%/FPS%/
    K%/BB%/Whiff% come from the Master-tab-derived `overall` dict (see
    pitcher_master_stats) instead of the Totals sheet's S%/FPS%/K%/BB%/
    SW-M% columns. FB Velo and Avg EV Allowed were already Master-derived
    either way."""
    def pct(key, val):
        return bm.interpolate(val, bm.PITCHER[key]) if val is not None else None

    velo_ev = velo_ev or {}
    fb_velo = _num(velo_ev.get('fb_velo'))
    avg_ev_allowed = _num(velo_ev.get('avg_ev_allowed'))
    overall = overall or {}

    rows = [
        {'label': 'FB Velo', 'value': fb_velo, 'pctl': pct('fb_velo', fb_velo), 'unit': ''},
        {'label': 'Avg EV Allowed', 'value': avg_ev_allowed, 'pctl': pct('avg_ev_allowed', avg_ev_allowed), 'unit': ''},
        {'label': 'Strike%', 'value': overall.get('strike_pct'), 'pctl': pct('strike_pct', overall.get('strike_pct')), 'unit': '%'},
        {'label': 'FPS%', 'value': overall.get('fps_pct'), 'pctl': pct('fps_pct', overall.get('fps_pct')), 'unit': '%'},
        {'label': 'K%', 'value': overall.get('k_pct'), 'pctl': pct('k_pct', overall.get('k_pct')), 'unit': '%'},
        {'label': 'BB%', 'value': overall.get('bb_pct'), 'pctl': pct('bb_pct', overall.get('bb_pct')), 'unit': '%'},
        {'label': 'Whiff%', 'value': overall.get('whiff_pct'), 'pctl': pct('whiff_pct', overall.get('whiff_pct')), 'unit': '%'},
    ]
    return [r for r in rows if r['value'] is not None]


def build_hitter_percentiles(totals_row, derived):
    def pct(key, val):
        return bm.interpolate(val, bm.HITTER[key]) if val is not None else None

    ev90 = _num(totals_row.get('EV90'))
    bs90 = _num(totals_row.get('BatSpeed90'))
    k_pct = _num(totals_row.get('K%'))
    k_pct = k_pct * 100 if k_pct is not None else None
    bb_pct = _num(totals_row.get('BB%'))
    bb_pct = bb_pct * 100 if bb_pct is not None else None
    whiff_pct = _num(totals_row.get('Whiff %'))
    whiff_pct = whiff_pct * 100 if whiff_pct is not None else None
    top50 = derived.get('top50_dist')
    angle = derived.get('avg_launch_angle')
    squp = derived.get('squared_up_pct')
    launch_q_pctl = pct('launch_quality', abs(angle - 18)) if angle is not None else None

    rows = [
        {'label': 'EV90', 'value': ev90, 'pctl': pct('ev90', ev90), 'unit': ''},
        {'label': 'Bat Speed 90', 'value': bs90, 'pctl': pct('bat_speed', bs90), 'unit': ''},
        {'label': 'Top 50% Dist', 'value': top50, 'pctl': pct('top50_dist', top50), 'unit': 'ft'},
        {'label': 'Avg Launch °', 'value': angle, 'pctl': launch_q_pctl, 'unit': '°'},
        {'label': 'Squared Up %', 'value': squp, 'pctl': pct('squared_up', squp), 'unit': '%'},
        {'label': 'Whiff %', 'value': whiff_pct, 'pctl': pct('whiff_pct', whiff_pct), 'unit': '%'},
        {'label': 'K%', 'value': k_pct, 'pctl': pct('k_pct', k_pct), 'unit': '%'},
        {'label': 'BB%', 'value': bb_pct, 'pctl': pct('bb_pct', bb_pct), 'unit': '%'},
    ]
    return [r for r in rows if r['value'] is not None]


# --- Trackman (optional, pitching-only) --------------------------------
# Full Swing's own export has no break/spin data at all (see the Pitch
# Movement Profile / Spin Direction "sample values" warning notes), but some
# players also get a Trackman report from their school -- this section reads
# that CSV and fills those two charts with real numbers when one's attached.
#
# VERIFIED against a real export (LiveBpPitching, Oct 2026 -- a single
# live-BP bullpen session, "Last, First" name format):
#   - Column names: 'Pitcher', 'TaggedPitchType', 'HorzBreak',
#     'InducedVertBreak', 'SpinRate', 'SpinAxis' all matched exactly as
#     assumed. TaggedPitchType values like "ChangeUp" fold fine through the
#     existing lowercase/strip normalization.
#   - Pitcher name is "Last, First" in this export; the comma-swap in
#     _normalize_name() already handles it correctly.
#   - Trackman ALSO exports a ready-made clock-face string in a 'Tilt'
#     column (e.g. "1:45") -- no degree math needed when it's present, and
#     it's what's actually used below in preference to deriving one.
#   - Cross-checking this file's 'Tilt' against its own raw 'SpinAxis'
#     degrees disproved the original guess (0 deg = 12:00): the real mapping
#     is 180 deg = 12:00, increasing clockwise -- i.e.
#     clock_minutes = ((SpinAxis_deg - 180) % 360) * 2. Fixed below; this is
#     now only a fallback for a file that has SpinAxis but no Tilt column.
#   - This export also carries a Hawk-Eye-style 3D measurement block
#     (SpinAxis3dTransverseAngle/LongitudinalAngle/ActiveSpinRate/
#     SpinEfficiency/Tilt) -- a genuine second, *observed* axis distinct from
#     the inferred spin-based one, matching what Baseball Savant's own
#     spin-based-vs-observed chart actually compares. It was empty in every
#     row of this particular session (not every Trackman setup measures it),
#     so the code below reads it when populated and falls back to the
#     spin-based value (same as before) when it isn't.
#   - Induced Vertical Break is preferred over raw Vertical Break when both
#     are present, since IVB (gravity removed) is what the movement-profile
#     chart is meant to show -- unchanged, confirmed both columns coexist in
#     this export.

TRACKMAN_TYPE_MAP = {
    'fourseamfastball': 'Fastball', 'fourseam': 'Fastball', 'fastball': 'Fastball',
    'twoseamfastball': 'Fastball', 'twoseam': 'Fastball',
    'sinker': 'Sinker',
    'slider': 'Slider', 'sweeper': 'Slider',
    'changeup': 'Changeup', 'change': 'Changeup',
    'curveball': 'Curveball', 'curve': 'Curveball', 'knucklecurve': 'Curveball',
    'cutter': 'Cutter',
    'splitter': 'Splitter', 'splitfinger': 'Splitter',
}

# Finer labels for the Trackman-only mini report, matching Savant's own
# pitch-type chips (4-Seam / Sinker / Slider / Sweeper / Cutter / Split /
# Curve / Change). The main map above folds these into Full Swing's coarser
# labels so Trackman rows can merge into Full-Swing-derived tracking rows;
# the mini report has no Full Swing rows to line up with, so it keeps them.
TRACKMAN_FINE_TYPE_MAP = {
    'fourseamfastball': '4-Seam', 'fourseam': '4-Seam', 'fastball': '4-Seam', '4seam': '4-Seam',
    'twoseamfastball': 'Sinker', 'twoseam': 'Sinker', 'sinker': 'Sinker',
    'slider': 'Slider', 'sweeper': 'Sweeper',
    'changeup': 'Change', 'change': 'Change',
    'curveball': 'Curve', 'curve': 'Curve', 'knucklecurve': 'Curve',
    'cutter': 'Cutter',
    'splitter': 'Split', 'splitfinger': 'Split',
}
# Sampled from the client's reference pitch-type legend.
MINI_PITCH_COLORS = {
    'Sinker': '#f0a139', 'Curve': '#5fcee9', '4-Seam': '#c13d4d', 'Slider': '#ede750',
    'Cutter': '#894432', 'Split': '#5daaab', 'Sweeper': '#d6b552', 'Change': '#5abb4e',
}

_TM_PITCHER_COLS = ['Pitcher', 'PitcherName', 'Pitcher Name']
_TM_TYPE_COLS = ['TaggedPitchType', 'AutoPitchType', 'PitchType', 'Pitch Type']
_TM_HB_COLS = ['HorzBreak', 'Horizontal Break', 'HBreak', 'HB']
_TM_IVB_COLS = ['InducedVertBreak', 'Induced Vertical Break', 'IVB']
_TM_VB_COLS = ['VertBreak', 'Vertical Break', 'VB']
_TM_SPIN_RATE_COLS = ['SpinRate', 'Spin Rate']
_TM_SPIN_AXIS_COLS = ['SpinAxis', 'Spin Axis']
# Ready-made clock-face string Trackman supplies directly (preferred -- no
# degree math, and it's quantized/rounded the same way Trackman's own
# reports show it).
_TM_TILT_COLS = ['Tilt']
# The true measured ("observed") axis, when this export's Trackman/Hawk-Eye
# setup actually captures it -- also a ready clock-face string.
_TM_OBS_TILT_COLS = ['SpinAxis3dTilt', 'Spin Axis 3d Tilt']
_TM_SPIN_EFF_COLS = ['SpinAxis3dSpinEfficiency', 'SpinEfficiency', 'Spin Efficiency', 'Spin Efficiency (release)']
_TM_VELO_COLS = ['RelSpeed', 'Velocity', 'Velo', 'PitchVelocity', 'Pitch Velocity']
_TM_THROWS_COLS = ['PitcherThrows', 'Pitcher Throws']
_TM_TEAM_COLS = ['PitcherTeam', 'Pitcher Team']
_TM_DATE_COLS = ['Date']
_TM_LOC_SIDE_COLS = ['PlateLocSide', 'Plate Loc Side']
_TM_LOC_HEIGHT_COLS = ['PlateLocHeight', 'Plate Loc Height']
_TM_CALL_COLS = ['PitchCall', 'Pitch Call']
_TM_BATSIDE_COLS = ['BatterSide', 'Batter Side']
_TM_EV_COLS = ['ExitSpeed', 'Exit Speed', 'ExitVelocity']
_TM_VAA_COLS = ['VertApprAngle', 'Vert Appr Angle', 'VerticalApproachAngle']
# a pitch counts as a strike for Strike% if it was called/swinging, fouled, or put in play
_TM_STRIKE_CALLS = {'strikecalled', 'strikeswinging', 'foulball', 'foulballnotfieldable', 'foulballfieldable', 'inplay'}
_TM_BALL_CALLS = {'ballcalled', 'ballintheDirt'.lower(), 'ballinDirt'.lower(), 'hitbypitch', 'intentionalball'}
_TM_BALLS_COLS = ['Balls']
_TM_STRIKES_COLS = ['Strikes']


def _tm_get(row, candidates):
    """Case/space/dash-insensitive header lookup against whatever columns
    this particular Trackman export actually has."""
    for c in candidates:
        if c in row:
            return row[c]
    norm_targets = {c.strip().lower().replace(' ', '').replace('-', '') for c in candidates}
    for k, v in row.items():
        nk = (k or '').strip().lower().replace(' ', '').replace('-', '')
        if nk in norm_targets:
            return v
    return None


def _tm_num(v):
    try:
        if v is None or str(v).strip() == '':
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def _normalize_name(raw):
    if not raw:
        return ''
    raw = str(raw).strip()
    if ',' in raw:
        last, first = [p.strip() for p in raw.split(',', 1)]
        raw = f'{first} {last}'
    return ' '.join(sorted(re.findall(r"[a-z]+", raw.lower())))


def _normalize_pitch_type(raw, fine=False):
    if not raw:
        return None
    key = str(raw).strip().lower().replace(' ', '').replace('-', '')
    return (TRACKMAN_FINE_TYPE_MAP if fine else TRACKMAN_TYPE_MAP).get(key)


def _circular_mean_minutes(minutes_list):
    """Circular mean over clock-face positions expressed as minutes past
    12:00 on a 12-hour face (0-720), so averaging across the 12:00/0:00
    wrap-around (e.g. 11:45 and 12:15) comes out right instead of landing on
    the wrong side of the clock."""
    if not minutes_list:
        return None
    angles = [m / 720 * 2 * math.pi for m in minutes_list]
    sin_sum = sum(math.sin(a) for a in angles)
    cos_sum = sum(math.cos(a) for a in angles)
    mean_angle = math.atan2(sin_sum, cos_sum) % (2 * math.pi)
    return mean_angle / (2 * math.pi) * 720


def _minutes_to_clock(total_minutes, quantize_to=15):
    """Minutes-past-12:00 -> 'H:MM' string. Rounds to the nearest 15-minute
    mark by default, matching how Trackman's own Tilt column is quantized
    (pitch-to-pitch spin-axis wobble makes finer precision look falsely
    exact)."""
    if total_minutes is None:
        return None
    if quantize_to:
        total_minutes = round(total_minutes / quantize_to) * quantize_to
    total_minutes = total_minutes % 720
    hour = int(total_minutes // 60) % 12
    minute = int(round(total_minutes % 60)) % 60
    return f'{hour or 12}:{minute:02d}'


def _spinaxis_deg_to_minutes(deg):
    """Converts TrackMan's raw SpinAxis degree value to the same clock-face
    'minutes past 12:00' convention as its own Tilt column -- verified
    against a real export (SpinAxis=231.78 -> Tilt=1:45, 239.79 -> 2:00,
    213.36 -> 1:00, etc.): 12:00 sits at SpinAxis=180 degrees, increasing
    clockwise from there. Only used as a fallback when a file has SpinAxis
    but no ready-made Tilt column."""
    if deg is None:
        return None
    return ((deg - 180) % 360) * 2  # 360 degrees <-> 720 clock-minutes


def _clock_str_to_minutes(s):
    """Parses a 'H:MM' clock-face string (as Trackman's Tilt/SpinAxis3dTilt
    columns give it) into minutes past 12:00 on a 12-hour face."""
    if not s:
        return None
    m = re.match(r'^\s*(\d{1,2}):(\d{2})\s*$', str(s))
    if not m:
        return None
    hour, minute = int(m.group(1)), int(m.group(2))
    return (hour % 12) * 60 + minute


def _gyro_deg_from_efficiency(eff):
    if eff is None:
        return None
    eff = eff / 100 if eff > 1.5 else eff  # tolerate either 0-1 or 0-100 scale
    eff = max(-1.0, min(1.0, eff))
    return round(math.degrees(math.acos(eff)))


def load_trackman_rows(csv_path):
    """Reads a Trackman per-pitch CSV export into a list of plain dicts.
    No fixed schema assumed beyond "it has headers" -- pitcher_trackman_movement()
    does its own flexible column matching against whatever's actually there.
    Missing/not-yet-uploaded file just means no Trackman data -- same as a
    player who never got one -- rather than an error."""
    if not csv_path or not os.path.exists(csv_path):
        return []
    with open(csv_path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


def pitcher_trackman_movement(trackman_rows, pitcher_name, fine=False):
    """Groups a Trackman export's break/spin numbers by pitch type for one
    pitcher, normalized to this app's own pitch-type labels (Fastball,
    Slider, Changeup, Curveball, Cutter, Splitter, Sinker) so they line up
    with the Full-Swing-derived `tracking` rows they're merged into.
    Returns {type: {'hb', 'ivb', 'spin_rate', 'spin_based_clock',
    'observed_clock', 'spin_deviation'}}."""
    if not trackman_rows:
        return {}
    target = _normalize_name(pitcher_name)
    by_type = defaultdict(lambda: {
        'hb': [], 'ivb': [], 'spin_rate': [],
        'tilt_min': [], 'obs_tilt_min': [], 'spin_eff': [], 'points': [],
    })

    for row in trackman_rows:
        raw_name = _tm_get(row, _TM_PITCHER_COLS)
        if _normalize_name(raw_name) != target:
            continue
        label = _normalize_pitch_type(_tm_get(row, _TM_TYPE_COLS), fine=fine)
        if not label:
            continue
        d = by_type[label]
        hb = _tm_num(_tm_get(row, _TM_HB_COLS))
        ivb = _tm_num(_tm_get(row, _TM_IVB_COLS))
        if ivb is None:
            ivb = _tm_num(_tm_get(row, _TM_VB_COLS))  # raw break as a fallback, not induced
        spin_rate = _tm_num(_tm_get(row, _TM_SPIN_RATE_COLS))
        spin_eff = _tm_num(_tm_get(row, _TM_SPIN_EFF_COLS))

        # Prefer Trackman's own ready-made clock string; fall back to
        # deriving one from the raw SpinAxis degrees only if Tilt is absent.
        tilt_min = _clock_str_to_minutes(_tm_get(row, _TM_TILT_COLS))
        if tilt_min is None:
            tilt_min = _spinaxis_deg_to_minutes(_tm_num(_tm_get(row, _TM_SPIN_AXIS_COLS)))
        obs_tilt_min = _clock_str_to_minutes(_tm_get(row, _TM_OBS_TILT_COLS))

        if hb is not None and ivb is not None:
            d['points'].append((hb, ivb, _tm_num(_tm_get(row, _TM_VELO_COLS))))
        if hb is not None:
            d['hb'].append(hb)
        if ivb is not None:
            d['ivb'].append(ivb)
        if spin_rate is not None:
            d['spin_rate'].append(spin_rate)
        if tilt_min is not None:
            d['tilt_min'].append(tilt_min)
        if obs_tilt_min is not None:
            d['obs_tilt_min'].append(obs_tilt_min)
        if spin_eff is not None:
            d['spin_eff'].append(spin_eff)

    out = {}
    for label, d in by_type.items():
        spin_based_clock = _minutes_to_clock(_circular_mean_minutes(d['tilt_min']))
        # No real "observed" (Hawk-Eye 3D) measurement for this pitch type ->
        # same simplification as before: show the spin-based value in both.
        observed_clock = _minutes_to_clock(_circular_mean_minutes(d['obs_tilt_min'])) if d['obs_tilt_min'] else spin_based_clock
        out[label] = {
            'points': d['points'],
            'hb': round(sum(d['hb']) / len(d['hb']), 1) if d['hb'] else None,
            'ivb': round(sum(d['ivb']) / len(d['ivb']), 1) if d['ivb'] else None,
            'spin_rate': round(sum(d['spin_rate']) / len(d['spin_rate'])) if d['spin_rate'] else None,
            'spin_based_clock': spin_based_clock,
            'observed_clock': observed_clock,
            'spin_deviation': _gyro_deg_from_efficiency(sum(d['spin_eff']) / len(d['spin_eff'])) if d['spin_eff'] else None,
        }
    return out


def pitcher_trackman_summary(trackman_rows, pitcher_name):
    """Builds a Trackman-only pitch-by-pitch summary for a pitcher who has
    no Full Swing workbook yet -- usage%, velo (avg/top) and the
    movement/spin numbers from pitcher_trackman_movement(), all computed
    straight from this one CSV. Deliberately has no season line, percentiles
    or comps -- those need a Full Swing export this player doesn't have.
    Returns (meta, pitches): meta = {'throws','team','date','n_pitches'},
    pitches = per-type dicts sorted by usage desc."""
    if not trackman_rows:
        return {}, []
    target = _normalize_name(pitcher_name)
    movement = pitcher_trackman_movement(trackman_rows, pitcher_name, fine=True)

    by_type = defaultdict(lambda: {'n': 0, 'velo': [], 'points': [], 'loc': [], 'spin': [], 'vaa': [], 'strikes': 0, 'called': 0})
    meta = {'throws': None, 'team': None, 'date': None, 'n_pitches': 0, 'log': [], 'release': []}
    for row in trackman_rows:
        raw_name = _tm_get(row, _TM_PITCHER_COLS)
        if _normalize_name(raw_name) != target:
            continue
        meta['n_pitches'] += 1
        if meta['throws'] is None:
            meta['throws'] = _tm_get(row, _TM_THROWS_COLS)
        if meta['team'] is None:
            meta['team'] = _tm_get(row, _TM_TEAM_COLS)
        if meta['date'] is None:
            meta['date'] = _tm_get(row, _TM_DATE_COLS)
        label = _normalize_pitch_type(_tm_get(row, _TM_TYPE_COLS), fine=True)
        if not label:
            continue
        d = by_type[label]
        d['n'] += 1
        velo = _tm_num(_tm_get(row, _TM_VELO_COLS))
        _rh = _tm_num(_tm_get(row, ['RelHeight', 'Release Height']))
        _rs = _tm_num(_tm_get(row, ['RelSide', 'Release Side']))
        if _rh is not None and _rs is not None:
            meta['release'].append({'type': label, 'h': _rh, 's': _rs,
                                    'ext': _tm_num(_tm_get(row, ['Extension', 'Release Extension']))})
        # --- pitch log row (one per tagged pitch, in outing order)
        _call = (_tm_get(row, _TM_CALL_COLS) or '').strip()
        _batter = (_tm_get(row, ['Batter', 'BatterName', 'Batter Name']) or '').strip()
        if ',' in _batter:
            _last, _first = [x.strip() for x in _batter.split(',', 1)]
            _batter = f'{_first} {_last}'.strip()
        _tilt = _tm_get(row, _TM_TILT_COLS)
        _tmin = _clock_str_to_minutes(_tilt)
        if _tmin is None:
            _ax = _tm_num(_tm_get(row, _TM_SPIN_AXIS_COLS))
            _tmin = _spinaxis_deg_to_minutes(_ax) if _ax is not None else None
        _ivb_l = _tm_num(_tm_get(row, _TM_IVB_COLS))
        if _ivb_l is None:
            _ivb_l = _tm_num(_tm_get(row, _TM_VB_COLS))
        meta['log'].append({
            'n': len(meta['log']) + 1, 'batter': _batter,
            'bats': (_tm_get(row, _TM_BATSIDE_COLS) or '').strip()[:1].upper(),
            'type': label, 'velo': _tm_num(_tm_get(row, _TM_VELO_COLS)),
            'ivb': _ivb_l, 'hb': _tm_num(_tm_get(row, _TM_HB_COLS)),
            'spin': _tm_num(_tm_get(row, _TM_SPIN_RATE_COLS)),
            'spin_dir': _minutes_to_clock(_tmin) if _tmin is not None else None,
            'ext': _tm_num(_tm_get(row, ['Extension', 'Release Extension'])),
            'vaa': _tm_num(_tm_get(row, _TM_VAA_COLS)),
            'ev': _tm_num(_tm_get(row, _TM_EV_COLS)),
            'balls': _tm_num(_tm_get(row, _TM_BALLS_COLS)),
            'call': _call,
        })
        if velo is not None:
            d['velo'].append(velo)
        spin = _tm_num(_tm_get(row, _TM_SPIN_RATE_COLS))
        if spin is not None:
            d['spin'].append(spin)
        vaa = _tm_num(_tm_get(row, _TM_VAA_COLS))
        if vaa is not None:
            d['vaa'].append(vaa)
        call_key = (_tm_get(row, _TM_CALL_COLS) or '').strip().lower()
        if call_key in _TM_STRIKE_CALLS:
            d['strikes'] += 1; d['called'] += 1
        elif call_key in _TM_BALL_CALLS:
            d['called'] += 1
        # every individual pitch, for the movement-profile scatter
        hb = _tm_num(_tm_get(row, _TM_HB_COLS))
        ivb = _tm_num(_tm_get(row, _TM_IVB_COLS))
        if ivb is None:
            ivb = _tm_num(_tm_get(row, _TM_VB_COLS))
        if hb is not None and ivb is not None:
            d['points'].append((hb, ivb, velo))
        # plate location (feet; side + = catcher's right / 1B side), for the zone heat map
        side = _tm_num(_tm_get(row, _TM_LOC_SIDE_COLS))
        height = _tm_num(_tm_get(row, _TM_LOC_HEIGHT_COLS))
        if side is not None and height is not None:
            d['loc'].append({
                'side': side, 'height': height, 'velo': velo, 'n': len(meta['log']),
                'call': (_tm_get(row, _TM_CALL_COLS) or '').strip(),
                'bats': (_tm_get(row, _TM_BATSIDE_COLS) or '').strip().lower()[:1].upper(),  # 'L' / 'R' / ''
                'count': f"{(_tm_get(row, _TM_BALLS_COLS) or '').strip()}-{(_tm_get(row, _TM_STRIKES_COLS) or '').strip()}",
            })

    total = sum(d['n'] for d in by_type.values()) or 1
    pitches = []
    for label, d in by_type.items():
        tm = movement.get(label, {})
        pitches.append({
            'type': label,
            'n': d['n'],
            'pct': d['n'] / total * 100,
            'velo': round(sum(d['velo']) / len(d['velo']), 1) if d['velo'] else None,
            'top_velo': round(max(d['velo']), 1) if d['velo'] else None,
            'points': d['points'], 'loc': d['loc'],
            'strike_pct': (d['strikes'] / d['called'] * 100) if d['called'] else None,
            'spin_avg': round(sum(d['spin']) / len(d['spin'])) if d['spin'] else None,
            'spin_max': round(max(d['spin'])) if d['spin'] else None,
            'vaa': round(sum(d['vaa']) / len(d['vaa']), 1) if d['vaa'] else None,
            'hb': tm.get('hb'), 'ivb': tm.get('ivb'),
            'spin_rate': tm.get('spin_rate'),
            'spin_based_clock': tm.get('spin_based_clock'),
            'observed_clock': tm.get('observed_clock'),
            'spin_deviation': tm.get('spin_deviation'),
        })
    pitches.sort(key=lambda p: p['pct'], reverse=True)
    return meta, pitches


# Arm-slot model (frontal plane, as seen from behind the plate):
#   shoulder height  = 0.81 x standing height  (acromion height, standard anthropometric ratio)
#   shoulder offset  = 0.13 x standing height out from the torso midline toward the throwing arm
#                      (half of shoulder breadth), torso midline taken as the centre of the rubber
#   arm angle        = angle of the line shoulder -> release point above horizontal
#                      (0 = sidearm, 90 = straight over the top)
SHOULDER_HEIGHT_FRAC = 0.81
SHOULDER_OFFSET_FRAC = 0.13


def arm_slot_label(angle):
    if angle is None:
        return None
    if angle >= 70:
        return 'Over the Top'
    if angle >= 50:
        return 'High 3/4'
    if angle >= 35:
        return '3/4'
    if angle >= 20:
        return 'Low 3/4'
    if angle >= 5:
        return 'Sidearm'
    return 'Submarine'


def arm_slot_stats(release, height_in):
    """release: [{'type','h','s','ext'}] (feet, from the CSV); height_in: pitcher's standing
    height in inches. Returns None without both. Angles are computed per pitch, then averaged."""
    if not release or not height_in:
        return None
    H = height_in / 12.0
    mean_side = sum(r['s'] for r in release) / len(release)
    sgn = 1 if mean_side >= 0 else -1            # +1: arm-side release is on the positive RelSide
    sh_y = SHOULDER_HEIGHT_FRAC * H
    sh_x = SHOULDER_OFFSET_FRAC * H                 # distance out from midline (always positive)
    angles = []
    for r in release:
        reach = r['s'] * sgn - sh_x                 # horizontal distance shoulder -> hand
        a = math.degrees(math.atan2(r['h'] - sh_y, reach)) if reach > 0.05 else 90.0
        angles.append(max(-30.0, min(100.0, a)))
    avg_h = sum(r['h'] for r in release) / len(release)
    avg_s = sum(abs(r['s']) for r in release) / len(release)
    exts = [r['ext'] for r in release if r.get('ext') is not None]
    ang = sum(angles) / len(angles)
    return {'angle': ang, 'label': arm_slot_label(ang), 'rel_height': avg_h, 'rel_side': avg_s,
            'extension': (sum(exts) / len(exts)) if exts else None, 'n': len(angles),
            'sd': (sum((a - ang) ** 2 for a in angles) / len(angles)) ** 0.5,
            'height_in': height_in, 'shoulder_y': sh_y, 'shoulder_x': sh_x, 'sgn': sgn,
            'points': [{'type': r['type'], 'h': r['h'], 's': r['s']} for r in release]}


def build_running_metrics(sixty_yd_time):
    """60 Yard Dash time (seconds, hand- or laser-timed) -> both the raw time
    and a derived Sprint Speed in ft/s (60 yards = 180 ft, so ft/s = 180 /
    time), each with its own percentile. Kept separate from
    build_hitter_percentiles() -- this never feeds the main percentile board
    or the MLB-comp matching, only the dedicated Running panel."""
    sixty_yd_time = _num(sixty_yd_time)
    if sixty_yd_time is None or sixty_yd_time <= 0:
        return []
    sprint_speed = round(180 / sixty_yd_time, 1)
    rows = [
        {'label': '60 Yard Dash', 'value': sixty_yd_time,
         'pctl': bm.interpolate(sixty_yd_time, bm.HITTER['sixty_yd']), 'unit': 's'},
        {'label': 'Sprint Speed', 'value': sprint_speed,
         'pctl': bm.interpolate(sprint_speed, bm.HITTER['sprint_speed']), 'unit': ' ft/s'},
    ]
    return rows
